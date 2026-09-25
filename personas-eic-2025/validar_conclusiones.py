"""Auditoría independiente de las cinco conclusiones del informe.

Uso: python validar_conclusiones.py --zip /ruta/eic2025_micro_20_csv.zip \
    --tabulado /ruta/eic2025_micro_105.xlsx
"""
from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def official_row(path: Path) -> dict:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    it = sheet.iter_rows(values_only=True)
    columns = next(it)
    for row in it:
        if str(row[0]) == "203180000" and row[7] == "Valor":
            return dict(zip(columns, row))
    raise ValueError("No se encontró el municipio 203180000, estimador Valor")


def independent_people(path: Path) -> pd.DataFrame:
    frames = []
    columns = ["CVEGEO", "FACTOR", "EDAD", "SEXO", "PERTE_INDIGENA", "HLENGUA",
               "DHSERSAL1", "SERSALUD", "CONACT", "SITTRA", "SERVICIO_MEDICO",
               "ACTIVIDADES_C"]
    with zipfile.ZipFile(path) as archive, archive.open("personas20.csv") as stream:
        for chunk in pd.read_csv(stream, dtype=str, keep_default_na=False,
                                 usecols=columns, chunksize=125_000):
            selection = chunk.loc[chunk["CVEGEO"] == "20318"]
            if not selection.empty:
                frames.append(selection)
    data = pd.concat(frames, ignore_index=True)
    data["FACTOR"] = data["FACTOR"].astype(int)
    data["EDAD"] = pd.to_numeric(data["EDAD"], errors="coerce")
    return data


def audit(data: pd.DataFrame, ref: dict) -> dict:
    w = data["FACTOR"]
    age = data["EDAD"]
    all_people = int(w.sum())
    def count(mask):
        return int(w[mask].sum())
    def ratio(mask, domain):
        return round(100 * count(mask & domain) / count(domain), 2)

    occupied = data["CONACT"].isin(["10", "13", "14", "15", "16", "17", "18", "19", "20"])
    seeking = data["CONACT"] == "30"
    labor = occupied | seeking
    adult = age.between(12, 130)
    women = data["SEXO"] == "3"
    men = data["SEXO"] == "1"
    salary = occupied & data["SITTRA"].isin(["1", "2", "3"])
    sector = data["ACTIVIDADES_C"].str[:2]
    sectors = {c: (count(occupied & sector.eq(c)), ratio(sector.eq(c), occupied))
               for c in ("46", "72", "23")}
    result = {
        "registros": len(data), "poblacion": all_people,
        "menores_15": [count(age.between(0,14)), ratio(age.between(0,14), age.between(0,130))],
        "65_mas": [count(age.between(65,130)), ratio(age.between(65,130), age.between(0,130))],
        "indigena": [count(data["PERTE_INDIGENA"] == "1"), ratio(data["PERTE_INDIGENA"] == "1", w.gt(0))],
        "habla_lengua_3mas": [count(data["HLENGUA"] == "1"), ratio(data["HLENGUA"] == "1", age.between(3,130))],
        "sin_afiliacion": [count(data["DHSERSAL1"] == "09"), ratio(data["DHSERSAL1"] == "09", w.gt(0))],
        "atencion_privada_todos": [count(data["SERSALUD"] == "07"), ratio(data["SERSALUD"] == "07", w.gt(0))],
        "atencion_privada_usuarios": [count(data["SERSALUD"] == "07"),
              ratio(data["SERSALUD"] == "07", data["SERSALUD"].isin([f"{i:02}" for i in range(1,10)]))],
        "participacion_mujeres_12mas": [count(labor & women), ratio(labor, adult & women)],
        "participacion_hombres_12mas": [count(labor & men), ratio(labor, adult & men)],
        "asalariados": count(salary),
        "servicio_medico_laboral": [count(salary & data["SERVICIO_MEDICO"].eq("5")),
                                    ratio(data["SERVICIO_MEDICO"].eq("5"), salary)],
        "ocupados": count(occupied),
        "sectores": sectors,
        "tres_sectores": [sum(v[0] for v in sectors.values()),
                           round(100 * sum(v[0] for v in sectors.values()) / count(occupied), 2)],
    }
    exact = {"poblacion": "POBTOT", "menores_15": "POB0_14", "65_mas": "POB65_MAS",
             "ocupados": "POCUPADA"}
    for our, official in exact.items():
        ours = result[our][0] if isinstance(result[our], list) else result[our]
        assert ours == int(ref[official]), (our, ours, ref[official])
    approximate = {"indigena": "PCN_POB_IND", "habla_lengua_3mas": "PCN_P3YM_HLI",
                   "sin_afiliacion": "PCN_PSINDER", "atencion_privada_usuarios": "PCN_PUSU_IPRIV"}
    for our, official in approximate.items():
        assert result[our][1] == float(ref[official]), (our, result[our], ref[official])
    assert ratio(labor, adult) == float(ref["PCN_PEA"])
    assert ratio(data["SITTRA"].isin(["1", "2", "3"]), occupied) == float(ref["PCN_POCUP_ASA"])
    assert result["tres_sectores"] == [14312, 46.86]
    assert result["servicio_medico_laboral"] == [5450, 28.69]
    assert result["participacion_mujeres_12mas"] == [13450, 53.26]
    assert result["participacion_hombres_12mas"] == [17357, 77.72]
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", type=Path, required=True)
    parser.add_argument("--tabulado", type=Path, required=True)
    parser.add_argument("--salida", type=Path, default=Path("resultados/auditoria.json"))
    args = parser.parse_args()
    result = audit(independent_people(args.zip), official_row(args.tabulado))
    args.salida.parent.mkdir(parents=True, exist_ok=True)
    args.salida.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Auditoría correcta: {result['registros']} registros, {result['poblacion']} habitantes")
